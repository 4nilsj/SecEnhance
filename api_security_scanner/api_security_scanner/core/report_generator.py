"""
Multi-format report generator using Jinja2 templates.
Generates comprehensive security scan reports with findings from ZAP and custom plugins.
Supports HTML, PDF, Excel, XML, and JSON formats.
"""

import json
import xml.etree.ElementTree as ET
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Any, Optional
from jinja2 import Environment, FileSystemLoader, Template

from ..utils.logger import get_logger, LoggedTimer

# Optional imports for additional report formats
try:
    from reportlab.lib.pagesizes import letter, A4
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
    REPORTLAB_AVAILABLE = True
except ImportError:
    REPORTLAB_AVAILABLE = False

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

try:
    import weasyprint
    WEASYPRINT_AVAILABLE = True
except (ImportError, OSError):
    # WeasyPrint has issues on Windows, so we'll disable it
    WEASYPRINT_AVAILABLE = False


class ReportGenerator:
    """Generates HTML reports from scan results."""
    
    def __init__(self, templates_dir: str = "templates"):
        self.logger = get_logger(__name__)
        self.templates_dir = Path(templates_dir)
        self.templates_dir.mkdir(exist_ok=True)
        
        # Setup Jinja2 environment
        self.jinja_env = Environment(
            loader=FileSystemLoader(str(self.templates_dir)),
            autoescape=True
        )
        
        # Create default template if it doesn't exist
        self._create_default_template()
    
    def _create_default_template(self):
        """Create default HTML template if it doesn't exist."""
        template_file = self.templates_dir / "report_template.html"
        
        if not template_file.exists():
            template_content = self._get_default_template()
            with open(template_file, 'w', encoding='utf-8') as f:
                f.write(template_content)
            self.logger.info("Created default report template")
    
    def _get_default_template(self) -> str:
        """Get the default HTML template content."""
        return """<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>API Security Scan Report - {{ scan_data.scan_id }}</title>
    <style>
        body {
            font-family: 'Segoe UI', Tahoma, Geneva, Verdana, sans-serif;
            margin: 0;
            padding: 20px;
            background-color: #f5f5f5;
            color: #333;
        }
        .container {
            max-width: 1200px;
            margin: 0 auto;
            background: white;
            border-radius: 8px;
            box-shadow: 0 2px 10px rgba(0,0,0,0.1);
            overflow: hidden;
        }
        .header {
            background: linear-gradient(135deg, #667eea 0%, #764ba2 100%);
            color: white;
            padding: 30px;
            text-align: center;
        }
        .header h1 {
            margin: 0;
            font-size: 2.5em;
            font-weight: 300;
        }
        .header p {
            margin: 10px 0 0 0;
            opacity: 0.9;
        }
        .content {
            padding: 30px;
        }
        .summary {
            display: grid;
            grid-template-columns: repeat(auto-fit, minmax(200px, 1fr));
            gap: 20px;
            margin-bottom: 30px;
        }
        .summary-card {
            background: #f8f9fa;
            padding: 20px;
            border-radius: 8px;
            text-align: center;
            border-left: 4px solid #007bff;
        }
        .summary-card h3 {
            margin: 0 0 10px 0;
            color: #007bff;
            font-size: 2em;
        }
        .summary-card p {
            margin: 0;
            color: #666;
        }
        .section {
            margin-bottom: 40px;
        }
        .section h2 {
            color: #333;
            border-bottom: 2px solid #007bff;
            padding-bottom: 10px;
            margin-bottom: 20px;
        }
        .alert {
            background: white;
            border: 1px solid #ddd;
            border-radius: 8px;
            margin-bottom: 20px;
            overflow: hidden;
            box-shadow: 0 2px 4px rgba(0,0,0,0.1);
        }
        .alert-header {
            padding: 15px 20px;
            font-weight: bold;
            display: flex;
            justify-content: space-between;
            align-items: center;
        }
        .alert-high { background-color: #f8d7da; border-left: 4px solid #dc3545; }
        .alert-medium { background-color: #fff3cd; border-left: 4px solid #ffc107; }
        .alert-low { background-color: #d1ecf1; border-left: 4px solid #17a2b8; }
        .alert-info { background-color: #d4edda; border-left: 4px solid #28a745; }
        .alert-body {
            padding: 20px;
            border-top: 1px solid #ddd;
        }
        .alert-body p {
            margin: 10px 0;
        }
        .risk-badge {
            padding: 4px 12px;
            border-radius: 20px;
            font-size: 0.8em;
            font-weight: bold;
            text-transform: uppercase;
        }
        .risk-high { background-color: #dc3545; color: white; }
        .risk-medium { background-color: #ffc107; color: #212529; }
        .risk-low { background-color: #17a2b8; color: white; }
        .risk-info { background-color: #28a745; color: white; }
        .performance-table {
            width: 100%;
            border-collapse: collapse;
            margin-top: 20px;
        }
        .performance-table th,
        .performance-table td {
            padding: 12px;
            text-align: left;
            border-bottom: 1px solid #ddd;
        }
        .performance-table th {
            background-color: #f8f9fa;
            font-weight: bold;
        }
        .performance-table tr:hover {
            background-color: #f5f5f5;
        }
        .footer {
            background: #f8f9fa;
            padding: 20px;
            text-align: center;
            color: #666;
            border-top: 1px solid #ddd;
        }
        .code {
            background: #f8f9fa;
            padding: 10px;
            border-radius: 4px;
            font-family: 'Courier New', monospace;
            font-size: 0.9em;
            overflow-x: auto;
            border: 1px solid #e9ecef;
            white-space: pre-wrap;
            word-break: break-all;
        }
        .request-block {
            background: #e3f2fd;
            border-left: 4px solid #2196f3;
            margin: 10px 0;
        }
        .response-block {
            background: #f3e5f5;
            border-left: 4px solid #9c27b0;
            margin: 10px 0;
        }
        .headers-block {
            background: #e8f5e8;
            border-left: 4px solid #4caf50;
            margin: 10px 0;
        }
        .vulnerable-text {
            background: #ffebee;
            border-left: 4px solid #f44336;
            margin: 10px 0;
        }
        .url-highlight {
            background: #fff3e0;
            border-left: 4px solid #ff9800;
            margin: 10px 0;
        }
        .method-highlight {
            background: #e1f5fe;
            border-left: 4px solid #00bcd4;
            margin: 10px 0;
        }
        .parameter-highlight {
            background: #fce4ec;
            border-left: 4px solid #e91e63;
            margin: 10px 0;
        }
        .highlight {
            background-color: #ffeb3b;
            padding: 2px 4px;
            border-radius: 3px;
            font-weight: bold;
            color: #d32f2f;
        }
        .executive-summary {
            background: #f8f9fa;
            border: 1px solid #dee2e6;
            border-radius: 8px;
            padding: 20px;
            margin: 20px 0;
        }
        .summary-text h3 {
            color: #495057;
            margin-top: 20px;
            margin-bottom: 10px;
        }
        .summary-text ul {
            margin: 10px 0;
            padding-left: 20px;
        }
        .summary-text li {
            margin: 5px 0;
        }
        .risk-high {
            background: #f8d7da;
            border: 1px solid #f5c6cb;
            border-radius: 4px;
            padding: 10px;
            margin: 10px 0;
            color: #721c24;
        }
        .risk-medium {
            background: #fff3cd;
            border: 1px solid #ffeaa7;
            border-radius: 4px;
            padding: 10px;
            margin: 10px 0;
            color: #856404;
        }
        .risk-low {
            background: #d1ecf1;
            border: 1px solid #bee5eb;
            border-radius: 4px;
            padding: 10px;
            margin: 10px 0;
            color: #0c5460;
        }
        .issues-table-container {
            overflow-x: auto;
            margin: 20px 0;
        }
        .issues-table {
            width: 100%;
            border-collapse: collapse;
            background: white;
            border-radius: 8px;
            overflow: hidden;
            box-shadow: 0 2px 4px rgba(0,0,0,0.1);
        }
        .issues-table th {
            background: #343a40;
            color: white;
            padding: 12px;
            text-align: left;
            font-weight: 600;
        }
        .issues-table td {
            padding: 12px;
            border-bottom: 1px solid #dee2e6;
            vertical-align: top;
        }
        .issues-table tr:hover {
            background: #f8f9fa;
        }
        .issues-table tr.severity-high {
            border-left: 4px solid #dc3545;
        }
        .issues-table tr.severity-medium {
            border-left: 4px solid #ffc107;
        }
        .issues-table tr.severity-low {
            border-left: 4px solid #28a745;
        }
        .issues-table tr.severity-informational {
            border-left: 4px solid #17a2b8;
        }
        .status-open {
            background: #dc3545;
            color: white;
            padding: 4px 8px;
            border-radius: 4px;
            font-size: 0.8em;
            font-weight: bold;
        }
        .vuln-link {
            color: #007bff;
            text-decoration: none;
            font-weight: bold;
        }
        .vuln-link:hover {
            color: #0056b3;
            text-decoration: underline;
        }
        .back-to-top {
            position: fixed;
            bottom: 20px;
            right: 20px;
            background: #007bff;
            color: white;
            padding: 10px 15px;
            border-radius: 5px;
            text-decoration: none;
            font-size: 14px;
            box-shadow: 0 2px 5px rgba(0,0,0,0.2);
            z-index: 1000;
        }
        .back-to-top:hover {
            background: #0056b3;
            color: white;
            text-decoration: none;
        }
        .vulnerability-section {
            scroll-margin-top: 20px;
        }
        .url {
            color: #007bff;
            word-break: break-all;
        }
        .method {
            display: inline-block;
            padding: 2px 8px;
            border-radius: 4px;
            font-size: 0.8em;
            font-weight: bold;
            margin-right: 10px;
        }
        .method-get { background-color: #28a745; color: white; }
        .method-post { background-color: #007bff; color: white; }
        .method-put { background-color: #ffc107; color: #212529; }
        .method-delete { background-color: #dc3545; color: white; }
        .method-patch { background-color: #6f42c1; color: white; }
        .poc-section {
            margin-top: 15px;
            border: 1px solid #ddd;
            border-radius: 4px;
            background-color: #f8f9fa;
        }
        .poc-section summary {
            padding: 10px 15px;
            background-color: #e9ecef;
            cursor: pointer;
            font-weight: bold;
            border-radius: 4px 4px 0 0;
        }
        .poc-section summary:hover {
            background-color: #dee2e6;
        }
        .poc-content {
            padding: 15px;
        }
        .poc-content h5 {
            margin: 10px 0 5px 0;
            color: #495057;
        }
        .references {
            margin-top: 15px;
            padding: 10px;
            background-color: #f8f9fa;
            border-radius: 4px;
        }
        .references h5 {
            margin: 0 0 10px 0;
            color: #495057;
        }
        .references ul {
            margin: 0;
            padding-left: 20px;
        }
        .references li {
            margin: 5px 0;
        }
        .references a {
            color: #007bff;
            text-decoration: none;
        }
        .references a:hover {
            text-decoration: underline;
        }
        .endpoints-list {
            background: #f8f9fa;
            border: 1px solid #dee2e6;
            border-radius: 8px;
            padding: 15px;
            margin: 15px 0;
        }
        .endpoint-item {
            display: flex;
            align-items: center;
            margin: 8px 0;
            padding: 8px;
            background: white;
            border-radius: 4px;
            border-left: 3px solid #007bff;
        }
        .endpoint-item .method {
            margin-right: 10px;
            min-width: 60px;
        }
        .endpoint-item .url {
            flex: 1;
            font-family: 'Courier New', monospace;
            color: #007bff;
            word-break: break-all;
        }
        .endpoint-item .folder {
            margin-left: 10px;
            color: #6c757d;
            font-size: 0.9em;
            font-style: italic;
        }
    </style>
</head>
<body>
    <div class="container">
        <div class="header">
            <h1>API Security Scan Report</h1>
            <p>Generated on {{ report_date }} | Scan ID: {{ scan_data.scan_id }}</p>
        </div>
        
        <div class="content">
            <!-- Executive Summary -->
            <div class="section">
                <h2>Executive Summary</h2>
                <div class="executive-summary">
                    <div class="summary-text">
                        <p>This security assessment was conducted on <strong>{{ scan_data.target_url }}</strong> to identify potential security vulnerabilities and misconfigurations. The scan was performed using automated security testing tools and custom security plugins.</p>
                        
                        <h3>Key Findings:</h3>
                        <ul>
                            <li><strong>Total Issues Identified:</strong> {{ all_vulnerabilities|length }}</li>
                            <li><strong>High Risk Issues:</strong> {{ risk_counts.High or 0 }}</li>
                            <li><strong>Medium Risk Issues:</strong> {{ risk_counts.Medium or 0 }}</li>
                            <li><strong>Low Risk Issues:</strong> {{ risk_counts.Low or 0 }}</li>
                            <li><strong>Informational Issues:</strong> {{ risk_counts.Informational or 0 }}</li>
                        </ul>
                        
                        {% if parsed_requests %}
                        <h3>API Endpoints Tested:</h3>
                        <p>The following {{ parsed_requests|length }} API endpoints were identified and tested during this security assessment:</p>
                        <div class="endpoints-list">
                            {% for request in parsed_requests %}
                            <div class="endpoint-item">
                                <span class="method method-{{ request.method.lower() }}">{{ request.method }}</span>
                                <span class="url">{{ request.url }}</span>
                                {% if request.folder %}
                                <span class="folder">({{ request.folder }})</span>
                                {% endif %}
                            </div>
                            {% endfor %}
                        </div>
                        {% endif %}
                        
                        <h3>Risk Assessment:</h3>
                        {% if (risk_counts.High or 0) > 0 %}
                        <div class="risk-high">
                            <strong>⚠️ HIGH RISK:</strong> {{ risk_counts.High or 0 }} critical vulnerabilities require immediate attention. These issues pose significant security risks and should be addressed as a priority.
                        </div>
                        {% endif %}
                        {% if (risk_counts.Medium or 0) > 0 %}
                        <div class="risk-medium">
                            <strong>🔶 MEDIUM RISK:</strong> {{ risk_counts.Medium or 0 }} vulnerabilities should be addressed in the next security update cycle. These issues could potentially be exploited under certain conditions.
                        </div>
                        {% endif %}
                        {% if (risk_counts.Low or 0) > 0 %}
                        <div class="risk-low">
                            <strong>🔷 LOW RISK:</strong> {{ risk_counts.Low or 0 }} issues are informational or pose minimal security risk but should be reviewed for best practices.
                        </div>
                        {% endif %}
                        
                        <h3>Recommendations:</h3>
                        <ul>
                            <li>Address all High and Medium risk vulnerabilities immediately</li>
                            <li>Implement proper security headers and CORS policies</li>
                            <li>Review and implement rate limiting mechanisms</li>
                            <li>Conduct regular security assessments</li>
                            <li>Implement a security monitoring and alerting system</li>
                        </ul>
                    </div>
                </div>
            </div>
            
            <!-- Issues Summary Table -->
            <div class="section">
                <h2>Issues Summary</h2>
                <div class="issues-table-container">
                    <table class="issues-table">
                        <thead>
                            <tr>
                                <th>Issue</th>
                                <th>Severity</th>
                                <th>CVSS Score</th>
                                <th>Impact</th>
                                <th>Source</th>
                                <th>Status</th>
                            </tr>
                        </thead>
                        <tbody>
                            {% for vuln in all_vulnerabilities %}
                            <tr class="severity-{{ vuln.risk.lower() }}">
                                <td>
                                    <a href="#vuln-{{ loop.index }}" class="vuln-link">
                                        <strong>{{ vuln.name }}</strong>
                                    </a>
                                    <br><small>{{ vuln.url }}</small>
                                </td>
                                <td>
                                    <span class="risk-badge risk-{{ vuln.risk.lower() }}">{{ vuln.risk }}</span>
                                </td>
                                <td>{{ vuln.cvss_score }}</td>
                                <td>
                                    {% if vuln.risk == 'High' %}
                                        Critical security vulnerability that could lead to data breach, unauthorized access, or system compromise
                                    {% elif vuln.risk == 'Medium' %}
                                        Moderate security risk that could be exploited under certain conditions
                                    {% elif vuln.risk == 'Low' %}
                                        Minor security issue or best practice violation
                                    {% else %}
                                        Informational finding for security awareness
                                    {% endif %}
                                </td>
                                <td>
                                    {% if vuln.plugin_name %}
                                        {{ vuln.plugin_name }}
                                    {% else %}
                                        ZAP Scanner
                                    {% endif %}
                                </td>
                                <td>
                                    <span class="status-open">Open</span>
                                </td>
                            </tr>
                            {% endfor %}
                        </tbody>
                    </table>
                </div>
            </div>
            
            <!-- Scan Summary -->
            <div class="section">
                <h2>Scan Summary</h2>
                <div class="summary">
                    <div class="summary-card">
                        <h3>{{ scan_data.target_url }}</h3>
                        <p>Target URL</p>
                    </div>
                    <div class="summary-card">
                        <h3>{{ scan_data.status }}</h3>
                        <p>Scan Status</p>
                    </div>
                    <div class="summary-card">
                        <h3>{{ "%.2f"|format(scan_data.total_duration or 0) }}s</h3>
                        <p>Total Duration</p>
                    </div>
                    <div class="summary-card">
                        <h3>{{ all_vulnerabilities|length }}</h3>
                        <p>Total Issues</p>
                    </div>
                    <div class="summary-card">
                        <h3>{{ risk_counts.High or 0 }}</h3>
                        <p>High Risk</p>
                    </div>
                </div>
            </div>
            
            <!-- Risk Summary -->
            <div class="section">
                <h2>Risk Summary</h2>
                <div class="summary">
                    <div class="summary-card">
                        <h3>{{ risk_counts.High or 0 }}</h3>
                        <p>High Risk</p>
                    </div>
                    <div class="summary-card">
                        <h3>{{ risk_counts.Medium or 0 }}</h3>
                        <p>Medium Risk</p>
                    </div>
                    <div class="summary-card">
                        <h3>{{ risk_counts.Low or 0 }}</h3>
                        <p>Low Risk</p>
                    </div>
                    <div class="summary-card">
                        <h3>{{ risk_counts.Informational or 0 }}</h3>
                        <p>Informational</p>
                    </div>
                </div>
            </div>
            
            <!-- Performance Statistics -->
            {% if performance_stats %}
            <div class="section">
                <h2>Performance Statistics</h2>
                <table class="performance-table">
                    <thead>
                        <tr>
                            <th>Phase</th>
                            <th>Duration</th>
                            <th>Details</th>
                        </tr>
                    </thead>
                    <tbody>
                        {% for stat in performance_stats %}
                        <tr>
                            <td>{{ stat.phase_name }}</td>
                            <td>{{ "%.2f"|format(stat.duration) }}s</td>
                            <td>{{ stat.details or '-' }}</td>
                        </tr>
                        {% endfor %}
                    </tbody>
                </table>
            </div>
            {% endif %}
            
            <!-- Vulnerabilities -->
            {% if vulnerabilities %}
            <div class="section">
                <h2>Security Vulnerabilities</h2>
                {% for vuln in vulnerabilities %}
                <div id="vuln-{{ loop.index }}" class="alert alert-{{ vuln.risk.lower() }} vulnerability-section">
                    <div class="alert-header">
                        <span>
                            <span class="method method-{{ vuln.request.split()[0].lower() if vuln.request else 'get' }}">{{ vuln.request.split()[0] if vuln.request else 'GET' }}</span>
                            <span class="url">{{ vuln.url }}</span>
                        </span>
                        <span class="risk-badge risk-{{ vuln.risk.lower() }}">{{ vuln.risk }}</span>
                    </div>
                    <div class="alert-body">
                        <h4>{{ vuln.name }}</h4>
                        <p><strong>CVSS Score:</strong> {{ vuln.cvss_score }}</p>
                        <p><strong>Source:</strong> {{ vuln.plugin_name or 'ZAP Scanner' }}</p>
                        {% if vuln.plugin_name %}
                        <p><strong>Plugin:</strong> {{ vuln.plugin_name }}</p>
                        {% endif %}
                        <p><strong>Description:</strong> {{ vuln.description }}</p>
                        {% if vuln.solution %}
                        <p><strong>Solution:</strong> {{ vuln.solution }}</p>
                        {% endif %}
                        {% if vuln.cwe_id %}
                        <p><strong>CWE ID:</strong> {{ vuln.cwe_id }}</p>
                        {% endif %}
                        {% if vuln.wasc_id %}
                        <p><strong>WASC ID:</strong> {{ vuln.wasc_id }}</p>
                        {% endif %}
                        {% if vuln.evidence %}
                        <p><strong>Evidence:</strong> {{ vuln.evidence }}</p>
                        {% endif %}
                        {% if vuln.parameter %}
                        <p><strong>Parameter:</strong> {{ vuln.parameter }}</p>
                        {% endif %}
                        
                        <!-- Collapsible Proof-of-Concept -->
                        <details class="poc-section">
                            <summary>Proof-of-Concept Evidence</summary>
                            <div class="poc-content">
                                {% if vuln.url %}
                                <h5>Target URL:</h5>
                                <div class="code url-highlight">{{ vuln.url }}</div>
                                {% endif %}
                                
                                {% if vuln.method %}
                                <h5>HTTP Method:</h5>
                                <div class="code method-highlight">{{ vuln.method }}</div>
                                {% endif %}
                                
                                <h5>Complete HTTP Request:</h5>
                                <div class="code request-block">
                                    {% if vuln.request %}
                                        {{ vuln.request|replace('\\n', '\n')|replace('\\r', '\r') }}
                                    {% else %}
                                        {{ vuln.method or 'GET' }} {{ vuln.url or 'N/A' }} HTTP/1.1
                                        {% if vuln.headers %}
                                        {% for header, value in vuln.headers.items() %}
                                        {{ header }}: {{ value }}
                                        {% endfor %}
                                        {% endif %}
                                        {% if vuln.body %}
                                        
                                        {{ vuln.body }}
                                        {% endif %}
                                    {% endif %}
                                </div>
                                
                                <h5>HTTP Response ({{ vuln.response_code or 'N/A' }}):</h5>
                                <div class="code response-block">
                                    {% if vuln.response %}
                                        {{ vuln.response|replace('\\n', '\n')|replace('\\r', '\r') }}
                                    {% else %}
                                        Response data not available
                                    {% endif %}
                                </div>
                                
                                {% if vuln.evidence %}
                                <h5>Vulnerable Text/Evidence:</h5>
                                <div class="code vulnerable-text">
                                    <span class="highlight">{{ vuln.evidence }}</span>
                                </div>
                                {% endif %}
                                
                                {% if vuln.parameter %}
                                <h5>Vulnerable Parameter:</h5>
                                <div class="code parameter-highlight">{{ vuln.parameter }}</div>
                                {% endif %}
                                
                                {% if vuln.response_headers %}
                                <h5>Response Headers:</h5>
                                <div class="code headers-block">
                                    {% for header, value in vuln.response_headers.items() %}
                                    {{ header }}: {{ value }}
                                    {% endfor %}
                                </div>
                                {% endif %}
                            </div>
                        </details>
                        
                        {% if vuln.references %}
                        <div class="references">
                            <h5>References:</h5>
                            <ul>
                                {% for ref in vuln.references %}
                                <li><a href="{{ ref }}" target="_blank">{{ ref }}</a></li>
                                {% endfor %}
                            </ul>
                        </div>
                        {% endif %}
                    </div>
                </div>
                {% endfor %}
            </div>
            {% endif %}
            
            <!-- Scan Details -->
            <div class="section">
                <h2>Scan Details</h2>
                <div class="code">
                    <strong>Scan ID:</strong> {{ scan_data.scan_id }}<br>
                    <strong>Target URL:</strong> {{ scan_data.target_url }}<br>
                    <strong>Start Time:</strong> {{ scan_data.start_time }}<br>
                    <strong>End Time:</strong> {{ scan_data.end_time or 'N/A' }}<br>
                    <strong>Input Type:</strong> {{ scan_data.input_type }}<br>
                    <strong>Input Source:</strong> {{ scan_data.input_source }}<br>
                    {% if scan_data.scan_mode %}
                    <strong>ZAP Scan Mode:</strong> {{ scan_data.scan_mode|title }}<br>
                    {% endif %}
                    {% if scan_data.auth_type %}
                    <strong>Authentication:</strong> {{ scan_data.auth_type }}<br>
                    {% endif %}
                    {% if scan_data.error_message %}
                    <strong>Error:</strong> {{ scan_data.error_message }}<br>
                    {% endif %}
                </div>
            </div>
        </div>
        
        <div class="footer">
            <p>Report generated by API Security Scanner on {{ report_date }}</p>
        </div>
        
        <!-- Back to Top Link -->
        <a href="#" class="back-to-top">↑ Back to Top</a>
    </div>
</body>
</html>"""
    
    def generate_report(self, scan_data: Dict[str, Any], vulnerabilities: List[Dict[str, Any]], 
                       performance_stats: List[Dict[str, Any]], output_path: str, 
                       parsed_requests: Optional[List[Dict[str, Any]]] = None) -> bool:
        """
        Generate HTML report from scan data.
        
        Args:
            scan_data: Scan metadata
            vulnerabilities: List of vulnerabilities with proof-of-concept data
            performance_stats: List of performance statistics
            output_path: Path to save the report
            parsed_requests: List of parsed API requests from collection
            
        Returns:
            True if report was generated successfully
        """
        try:
            with LoggedTimer(self.logger, "Report generation"):
                # Calculate risk counts
                risk_counts = self._calculate_risk_counts_from_vulnerabilities(vulnerabilities)
                
                # Prepare template data
                template_data = {
                    'scan_data': scan_data,
                    'vulnerabilities': vulnerabilities,
                    'all_vulnerabilities': vulnerabilities,  # For the issues table
                    'performance_stats': performance_stats,
                    'risk_counts': risk_counts,
                    'parsed_requests': parsed_requests or [],
                    'report_date': datetime.now().strftime('%Y-%m-%d %H:%M:%S')
                }
                
                # Load and render template
                template = self.jinja_env.get_template('report_template.html')
                html_content = template.render(**template_data)
                
                # Write report to file
                output_file = Path(output_path)
                output_file.parent.mkdir(parents=True, exist_ok=True)
                
                with open(output_file, 'w', encoding='utf-8') as f:
                    f.write(html_content)
                
                self.logger.info(f"Report generated successfully: {output_path}")
                return True
                
        except Exception as e:
            self.logger.error(f"Failed to generate report: {e}")
            return False
    
    def _calculate_risk_counts_from_vulnerabilities(self, vulnerabilities: List[Dict[str, Any]]) -> Dict[str, int]:
        """Calculate risk level counts from vulnerabilities."""
        risk_counts = {
            'High': 0,
            'Medium': 0,
            'Low': 0,
            'Informational': 0
        }
        
        # Count vulnerabilities by risk level
        for vuln in vulnerabilities:
            risk_level = self._get_vuln_attr(vuln, 'risk', 'Informational')
            if risk_level in risk_counts:
                risk_counts[risk_level] += 1
        
        return risk_counts
    
    def generate_json_report(self, scan_data: Dict[str, Any], vulnerabilities: List[Any], 
                           performance_stats: List[Dict[str, Any]], output_path: str) -> bool:
        """
        Generate JSON report from scan data.
        
        Args:
            scan_data: Scan metadata
            vulnerabilities: List of vulnerabilities with proof-of-concept data
            performance_stats: List of performance statistics
            output_path: Path to save the JSON report
            
        Returns:
            True if report was generated successfully
        """
        try:
            with LoggedTimer(self.logger, "JSON report generation"):
                # Convert Vulnerability objects to dictionaries for JSON serialization
                vuln_dicts = []
                for vuln in vulnerabilities:
                    if isinstance(vuln, dict):
                        vuln_dicts.append(vuln)
                    else:
                        # Convert Vulnerability object to dict
                        vuln_dict = {
                            'id': getattr(vuln, 'id', ''),
                            'name': getattr(vuln, 'name', ''),
                            'description': getattr(vuln, 'description', ''),
                            'risk': getattr(vuln, 'risk', ''),
                            'cvss_score': getattr(vuln, 'cvss_score', ''),
                            'solution': getattr(vuln, 'solution', ''),
                            'references': getattr(vuln, 'references', []),
                            'cwe_id': getattr(vuln, 'cwe_id', ''),
                            'wasc_id': getattr(vuln, 'wasc_id', ''),
                            'request': getattr(vuln, 'request', ''),
                            'response': getattr(vuln, 'response', ''),
                            'url': getattr(vuln, 'url', ''),
                            'parameter': getattr(vuln, 'parameter', ''),
                            'evidence': getattr(vuln, 'evidence', ''),
                            'scan_id': getattr(vuln, 'scan_id', ''),
                            'timestamp': str(getattr(vuln, 'timestamp', ''))
                        }
                        vuln_dicts.append(vuln_dict)
                
                # Prepare report data
                report_data = {
                    'scan_metadata': scan_data,
                    'vulnerabilities': vuln_dicts,
                    'performance_stats': performance_stats,
                    'risk_counts': self._calculate_risk_counts_from_vulnerabilities(vulnerabilities),
                    'generated_at': datetime.now().isoformat(),
                    'report_version': '2.0'
                }
                
                # Write JSON report
                output_file = Path(output_path)
                output_file.parent.mkdir(parents=True, exist_ok=True)
                
                with open(output_file, 'w', encoding='utf-8') as f:
                    json.dump(report_data, f, indent=2, default=str)
                
                self.logger.info(f"JSON report generated successfully: {output_path}")
                return True
                
        except Exception as e:
            self.logger.error(f"Failed to generate JSON report: {e}")
            return False
    
    def generate_summary_report(self, scan_data: Dict[str, Any], zap_alerts: List[Dict[str, Any]], 
                              custom_alerts: List[Dict[str, Any]], performance_stats: List[Dict[str, Any]]) -> str:
        """
        Generate a text summary report.
        
        Args:
            scan_data: Scan metadata
            zap_alerts: List of ZAP alerts
            custom_alerts: List of custom plugin alerts
            performance_stats: List of performance statistics
            
        Returns:
            Text summary report
        """
        try:
            # Combine all vulnerabilities and calculate risk counts
            all_vulnerabilities = zap_alerts + custom_alerts
            risk_counts = self._calculate_risk_counts_from_vulnerabilities(all_vulnerabilities)
            
            summary = []
            summary.append("=" * 60)
            summary.append("API SECURITY SCAN SUMMARY")
            summary.append("=" * 60)
            summary.append(f"Scan ID: {scan_data.get('scan_id', 'N/A')}")
            summary.append(f"Target: {scan_data.get('target_url', 'N/A')}")
            summary.append(f"Status: {scan_data.get('status', 'N/A')}")
            summary.append(f"Duration: {scan_data.get('total_duration', 0):.2f}s")
            summary.append("")
            
            summary.append("RISK SUMMARY:")
            summary.append(f"  High: {risk_counts['High']}")
            summary.append(f"  Medium: {risk_counts['Medium']}")
            summary.append(f"  Low: {risk_counts['Low']}")
            summary.append(f"  Informational: {risk_counts['Informational']}")
            summary.append("")
            
            summary.append("ALERTS SUMMARY:")
            summary.append(f"  ZAP Alerts: {len(zap_alerts)}")
            summary.append(f"  Custom Plugin Alerts: {len(custom_alerts)}")
            summary.append("")
            
            if performance_stats:
                summary.append("PERFORMANCE STATISTICS:")
                for stat in performance_stats:
                    summary.append(f"  {stat.get('phase_name', 'Unknown')}: {stat.get('duration', 0):.2f}s")
                summary.append("")
            
            if zap_alerts:
                summary.append("ZAP ALERTS:")
                for alert in zap_alerts:
                    summary.append(f"  [{alert.get('risk_level', 'Unknown')}] {alert.get('name', 'Unknown')}")
                    summary.append(f"    URL: {alert.get('url', 'N/A')}")
                    summary.append(f"    Method: {alert.get('method', 'N/A')}")
                    summary.append("")
            
            if custom_alerts:
                summary.append("CUSTOM PLUGIN ALERTS:")
                for alert in custom_alerts:
                    summary.append(f"  [{alert.get('risk', 'Unknown')}] {alert.get('name', 'Unknown')}")
                    summary.append(f"    Plugin: {alert.get('plugin_name', 'N/A')}")
                    summary.append(f"    URL: {alert.get('url', 'N/A')}")
                    summary.append("")
            
            summary.append("=" * 60)
            
            return "\n".join(summary)
            
        except Exception as e:
            self.logger.error(f"Failed to generate summary report: {e}")
            return f"Error generating summary: {e}"
    
    def generate_pdf_report(self, scan_data: Dict[str, Any], vulnerabilities: List[Any], 
                          performance_stats: List[Dict[str, Any]], output_path: str) -> bool:
        """
        Generate PDF report from scan data.
        
        Args:
            scan_data: Scan metadata and configuration
            vulnerabilities: List of vulnerability findings
            performance_stats: Performance timing data
            output_path: Output file path for PDF report
            
        Returns:
            bool: True if successful, False otherwise
        """
        if not REPORTLAB_AVAILABLE:
            self.logger.error("ReportLab not available. Install with: pip install reportlab")
            return False
        
        try:
            with LoggedTimer(self.logger, "PDF report generation"):
                doc = SimpleDocTemplate(output_path, pagesize=A4)
                story = []
                
                # Get styles
                styles = getSampleStyleSheet()
                title_style = ParagraphStyle(
                    'CustomTitle',
                    parent=styles['Heading1'],
                    fontSize=24,
                    spaceAfter=30,
                    alignment=TA_CENTER,
                    textColor=colors.darkblue
                )
                
                heading_style = ParagraphStyle(
                    'CustomHeading',
                    parent=styles['Heading2'],
                    fontSize=16,
                    spaceAfter=12,
                    textColor=colors.darkblue
                )
                
                # Title
                story.append(Paragraph("API Security Scan Report", title_style))
                story.append(Spacer(1, 12))
                
                # Scan Information
                story.append(Paragraph("Scan Information", heading_style))
                scan_info = [
                    ["Scan ID:", scan_data.get('scan_id', 'N/A')],
                    ["Target URL:", scan_data.get('target_url', 'N/A')],
                    ["Status:", scan_data.get('status', 'N/A')],
                    ["Start Time:", scan_data.get('start_time', 'N/A')],
                    ["End Time:", scan_data.get('end_time', 'N/A')],
                    ["Duration:", f"{scan_data.get('total_duration', 0):.2f}s"]
                ]
                
                # Add template, scan mode, and plugin information if available
                if scan_data.get('scan_mode'):
                    scan_info.append(["ZAP Scan Mode:", scan_data.get('scan_mode', '').title()])
                if scan_data.get('template_used'):
                    scan_info.append(["Scan Template:", scan_data.get('template_used')])
                if scan_data.get('plugins_used'):
                    scan_info.append(["Plugins Used:", scan_data.get('plugins_used')])
                
                scan_table = Table(scan_info, colWidths=[2*inch, 4*inch])
                scan_table.setStyle(TableStyle([
                    ('BACKGROUND', (0, 0), (0, -1), colors.lightgrey),
                    ('TEXTCOLOR', (0, 0), (-1, -1), colors.black),
                    ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                    ('FONTNAME', (0, 0), (-1, -1), 'Helvetica'),
                    ('FONTSIZE', (0, 0), (-1, -1), 10),
                    ('BOTTOMPADDING', (0, 0), (-1, -1), 12),
                    ('BACKGROUND', (1, 0), (1, -1), colors.beige),
                    ('GRID', (0, 0), (-1, -1), 1, colors.black)
                ]))
                story.append(scan_table)
                story.append(Spacer(1, 20))
                
                # Risk Summary
                risk_counts = self._calculate_risk_counts(vulnerabilities)
                story.append(Paragraph("Risk Summary", heading_style))
                risk_data = [
                    ["Risk Level", "Count"],
                    ["High", str(risk_counts['High'])],
                    ["Medium", str(risk_counts['Medium'])],
                    ["Low", str(risk_counts['Low'])],
                    ["Informational", str(risk_counts['Informational'])]
                ]
                
                risk_table = Table(risk_data, colWidths=[2*inch, 1*inch])
                risk_table.setStyle(TableStyle([
                    ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
                    ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                    ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                    ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                    ('FONTSIZE', (0, 0), (-1, -1), 10),
                    ('BOTTOMPADDING', (0, 0), (-1, -1), 12),
                    ('BACKGROUND', (0, 1), (-1, 1), colors.lightgrey),
                    ('BACKGROUND', (0, 3), (-1, 3), colors.lightgrey),
                    ('GRID', (0, 0), (-1, -1), 1, colors.black)
                ]))
                story.append(risk_table)
                story.append(Spacer(1, 20))
                
                # Vulnerabilities
                if vulnerabilities:
                    story.append(Paragraph("Vulnerability Details", heading_style))
                    
                    vuln_data = [["Risk", "CVSS", "Name", "URL", "Source"]]
                    for vuln in vulnerabilities[:20]:  # Limit to first 20 for PDF
                        risk = self._get_vuln_attr(vuln, 'risk', self._get_vuln_attr(vuln, 'risk_level', 'Unknown'))
                        cvss_score = self._get_vuln_attr(vuln, 'cvss_score', 'N/A')
                        name = str(self._get_vuln_attr(vuln, 'name', 'Unknown'))[:40]  # Truncate long names
                        url = str(self._get_vuln_attr(vuln, 'url', 'N/A'))[:35]  # Truncate long URLs
                        source = self._get_vuln_attr(vuln, 'plugin_name', 'ZAP Scanner')
                        
                        vuln_data.append([risk, cvss_score, name, url, source])
                    
                    vuln_table = Table(vuln_data, colWidths=[0.8*inch, 0.6*inch, 2.2*inch, 1.8*inch, 1.6*inch])
                    vuln_table.setStyle(TableStyle([
                        ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
                        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                        ('FONTSIZE', (0, 0), (-1, -1), 8),
                        ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
                        ('GRID', (0, 0), (-1, -1), 1, colors.black),
                        ('VALIGN', (0, 0), (-1, -1), 'TOP')
                    ]))
                    story.append(vuln_table)
                    story.append(Spacer(1, 20))
                
                # Performance Statistics
                if performance_stats:
                    story.append(Paragraph("Performance Statistics", heading_style))
                    perf_data = [["Phase", "Duration (s)"]]
                    for stat in performance_stats:
                        perf_data.append([
                            stat.get('phase_name', 'Unknown'),
                            f"{stat.get('duration', 0):.2f}"
                        ])
                    
                    perf_table = Table(perf_data, colWidths=[3*inch, 1*inch])
                    perf_table.setStyle(TableStyle([
                        ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
                        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                        ('FONTSIZE', (0, 0), (-1, -1), 10),
                        ('BOTTOMPADDING', (0, 0), (-1, -1), 12),
                        ('GRID', (0, 0), (-1, -1), 1, colors.black)
                    ]))
                    story.append(perf_table)
                
                # Build PDF
                doc.build(story)
                
                self.logger.info(f"PDF report generated successfully: {output_path}")
                return True
                
        except Exception as e:
            self.logger.error(f"Failed to generate PDF report: {e}")
            return False
    
    def generate_excel_report(self, scan_data: Dict[str, Any], vulnerabilities: List[Any], 
                            performance_stats: List[Dict[str, Any]], output_path: str) -> bool:
        """
        Generate Excel report from scan data.
        
        Args:
            scan_data: Scan metadata and configuration
            vulnerabilities: List of vulnerability findings
            performance_stats: Performance timing data
            output_path: Output file path for Excel report
            
        Returns:
            bool: True if successful, False otherwise
        """
        if not OPENPYXL_AVAILABLE:
            self.logger.error("OpenPyXL not available. Install with: pip install openpyxl")
            return False
        
        try:
            with LoggedTimer(self.logger, "Excel report generation"):
                wb = Workbook()
                # Remove default sheet if it exists and is not None
                default_sheet = wb.active
                if default_sheet is not None:
                    wb.remove(default_sheet)
                
                # Summary sheet
                summary_ws = wb.create_sheet("Summary")
                self._create_summary_sheet(summary_ws, scan_data, vulnerabilities, performance_stats)
                
                # Vulnerabilities sheet
                if vulnerabilities:
                    vuln_ws = wb.create_sheet("Vulnerabilities")
                    self._create_vulnerabilities_sheet(vuln_ws, vulnerabilities)
                
                # Performance sheet
                if performance_stats:
                    perf_ws = wb.create_sheet("Performance")
                    self._create_performance_sheet(perf_ws, performance_stats)
                
                # Save workbook
                wb.save(output_path)
                
                self.logger.info(f"Excel report generated successfully: {output_path}")
                return True
                
        except Exception as e:
            self.logger.error(f"Failed to generate Excel report: {e}")
            return False
    
    def generate_xml_report(self, scan_data: Dict[str, Any], vulnerabilities: List[Any], 
                          performance_stats: List[Dict[str, Any]], output_path: str) -> bool:
        """
        Generate XML report from scan data.
        
        Args:
            scan_data: Scan metadata and configuration
            vulnerabilities: List of vulnerability findings
            performance_stats: Performance timing data
            output_path: Output file path for XML report
            
        Returns:
            bool: True if successful, False otherwise
        """
        try:
            with LoggedTimer(self.logger, "XML report generation"):
                # Create root element
                root = ET.Element("security_scan_report")
                root.set("version", "1.0")
                root.set("generated_at", datetime.now().isoformat())
                
                # Scan information
                scan_info = ET.SubElement(root, "scan_information")
                for key, value in scan_data.items():
                    elem = ET.SubElement(scan_info, key.replace(' ', '_').lower())
                    elem.text = str(value)
                
                # Risk summary
                risk_counts = self._calculate_risk_counts(vulnerabilities)
                risk_summary = ET.SubElement(root, "risk_summary")
                for risk_level, count in risk_counts.items():
                    risk_elem = ET.SubElement(risk_summary, "risk_level")
                    risk_elem.set("level", risk_level)
                    risk_elem.set("count", str(count))
                
                # Vulnerabilities
                if vulnerabilities:
                    vulns_elem = ET.SubElement(root, "vulnerabilities")
                    for vuln in vulnerabilities:
                        vuln_elem = ET.SubElement(vulns_elem, "vulnerability")
                        
                        # Handle both dict and Vulnerability object
                        if isinstance(vuln, dict):
                            vuln_dict = vuln
                        else:
                            # Convert Vulnerability object to dict
                            vuln_dict = {
                                'id': getattr(vuln, 'id', ''),
                                'name': getattr(vuln, 'name', ''),
                                'description': getattr(vuln, 'description', ''),
                                'risk': getattr(vuln, 'risk', ''),
                                'cvss_score': getattr(vuln, 'cvss_score', ''),
                                'solution': getattr(vuln, 'solution', ''),
                                'url': getattr(vuln, 'url', ''),
                                'parameter': getattr(vuln, 'parameter', ''),
                                'evidence': getattr(vuln, 'evidence', ''),
                                'cwe_id': getattr(vuln, 'cwe_id', ''),
                                'wasc_id': getattr(vuln, 'wasc_id', ''),
                                'plugin_name': getattr(vuln, 'plugin_name', ''),
                                'scan_id': getattr(vuln, 'scan_id', ''),
                                'timestamp': str(getattr(vuln, 'timestamp', ''))
                            }
                        
                        for key, value in vuln_dict.items():
                            if isinstance(value, (str, int, float)):
                                sub_elem = ET.SubElement(vuln_elem, key.replace(' ', '_').lower())
                                sub_elem.text = str(value)
                
                # Performance statistics
                if performance_stats:
                    perf_elem = ET.SubElement(root, "performance_statistics")
                    for stat in performance_stats:
                        stat_elem = ET.SubElement(perf_elem, "statistic")
                        for key, value in stat.items():
                            if isinstance(value, (str, int, float)):
                                sub_elem = ET.SubElement(stat_elem, key.replace(' ', '_').lower())
                                sub_elem.text = str(value)
                
                # Write XML to file
                tree = ET.ElementTree(root)
                ET.indent(tree, space="  ", level=0)  # Pretty print
                tree.write(output_path, encoding='utf-8', xml_declaration=True)
                
                self.logger.info(f"XML report generated successfully: {output_path}")
                return True
                
        except Exception as e:
            self.logger.error(f"Failed to generate XML report: {e}")
            return False
    
    def _create_summary_sheet(self, ws, scan_data: Dict[str, Any], vulnerabilities: List[Dict[str, Any]], 
                             performance_stats: List[Dict[str, Any]]):
        """Create summary sheet for Excel report."""
        # Title
        ws['A1'] = "API Security Scan Report"
        ws['A1'].font = Font(size=16, bold=True)
        ws.merge_cells('A1:D1')
        
        # Scan information
        row = 3
        ws[f'A{row}'] = "Scan Information"
        ws[f'A{row}'].font = Font(bold=True)
        row += 1
        
        scan_info = [
            ("Scan ID", scan_data.get('scan_id', 'N/A')),
            ("Target URL", scan_data.get('target_url', 'N/A')),
            ("Status", scan_data.get('status', 'N/A')),
            ("Start Time", scan_data.get('start_time', 'N/A')),
            ("End Time", scan_data.get('end_time', 'N/A')),
            ("Duration", f"{scan_data.get('total_duration', 0):.2f}s")
        ]
        
        # Add template, scan mode, and plugin information if available
        if scan_data.get('scan_mode'):
            scan_info.append(("ZAP Scan Mode", scan_data.get('scan_mode', '').title()))
        if scan_data.get('template_used'):
            scan_info.append(("Scan Template", scan_data.get('template_used')))
        if scan_data.get('plugins_used'):
            scan_info.append(("Plugins Used", scan_data.get('plugins_used')))
        
        for label, value in scan_info:
            ws[f'A{row}'] = label
            ws[f'B{row}'] = value
            row += 1
        
        # Risk summary
        row += 1
        ws[f'A{row}'] = "Risk Summary"
        ws[f'A{row}'].font = Font(bold=True)
        row += 1
        
        risk_counts = self._calculate_risk_counts(vulnerabilities)
        ws[f'A{row}'] = "Risk Level"
        ws[f'B{row}'] = "Count"
        ws[f'A{row}'].font = Font(bold=True)
        ws[f'B{row}'].font = Font(bold=True)
        row += 1
        
        for risk_level, count in risk_counts.items():
            ws[f'A{row}'] = risk_level
            ws[f'B{row}'] = count
            row += 1
        
        # Auto-adjust column widths
        for column in ws.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width
    
    def _create_vulnerabilities_sheet(self, ws, vulnerabilities: List[Dict[str, Any]]):
        """Create vulnerabilities sheet for Excel report."""
        # Headers
        headers = ["Risk Level", "CVSS Score", "Name", "URL", "Method", "Plugin", "Description", "Solution"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
        
        # Data
        for row, vuln in enumerate(vulnerabilities, 2):
            ws.cell(row=row, column=1, value=self._get_vuln_attr(vuln, 'risk', self._get_vuln_attr(vuln, 'risk_level', 'Unknown')))
            ws.cell(row=row, column=2, value=self._get_vuln_attr(vuln, 'cvss_score', 'N/A'))
            ws.cell(row=row, column=3, value=self._get_vuln_attr(vuln, 'name', 'Unknown'))
            ws.cell(row=row, column=4, value=self._get_vuln_attr(vuln, 'url', 'N/A'))
            ws.cell(row=row, column=5, value=self._get_vuln_attr(vuln, 'method', 'N/A'))
            ws.cell(row=row, column=6, value=self._get_vuln_attr(vuln, 'plugin_name', 'ZAP Scanner'))
            ws.cell(row=row, column=7, value=self._get_vuln_attr(vuln, 'description', 'N/A'))
            ws.cell(row=row, column=8, value=self._get_vuln_attr(vuln, 'solution', 'N/A'))
        
        # Auto-adjust column widths
        for column in ws.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width
    
    def _create_performance_sheet(self, ws, performance_stats: List[Dict[str, Any]]):
        """Create performance sheet for Excel report."""
        # Headers
        headers = ["Phase", "Duration (s)", "Start Time", "End Time"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
        
        # Data
        for row, stat in enumerate(performance_stats, 2):
            ws.cell(row=row, column=1, value=stat.get('phase_name', 'Unknown'))
            ws.cell(row=row, column=2, value=stat.get('duration', 0))
            ws.cell(row=row, column=3, value=stat.get('start_time', 'N/A'))
            ws.cell(row=row, column=4, value=stat.get('end_time', 'N/A'))
        
        # Auto-adjust column widths
        for column in ws.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width
    
    def _get_vuln_attr(self, vuln: Any, attr: str, default: Any = None) -> Any:
        """Safely get attribute from vulnerability (handles both dict and object)."""
        if hasattr(vuln, attr):
            return getattr(vuln, attr)
        elif isinstance(vuln, dict):
            return vuln.get(attr, default)
        else:
            return default
    
    def _calculate_risk_counts(self, vulnerabilities: List[Any]) -> Dict[str, int]:
        """Calculate risk level counts from vulnerabilities."""
        risk_counts = {'High': 0, 'Medium': 0, 'Low': 0, 'Informational': 0}
        
        for vuln in vulnerabilities:
            risk = self._get_vuln_attr(vuln, 'risk', self._get_vuln_attr(vuln, 'risk_level', 'Informational'))
            
            if risk in risk_counts:
                risk_counts[risk] += 1
            else:
                risk_counts['Informational'] += 1
        
        return risk_counts
