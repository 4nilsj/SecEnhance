#!/usr/bin/env python3
"""
Sample Data Creator for XLOOKUP Tool
Creates a sample Excel file with multiple sheets for testing.
"""

import openpyxl
from openpyxl.styles import Font, PatternFill
import random

def create_sample_excel():
    """Create a sample Excel file with multiple sheets."""
    
    # Create a new workbook
    wb = openpyxl.Workbook()
    
    # Remove default sheet
    wb.remove(wb.active)
    
    # Create Sales sheet
    sales_ws = wb.create_sheet("Sales")
    
    # Sales data headers
    sales_headers = ["ProductID", "CustomerID", "Quantity", "Price", "Date"]
    for col, header in enumerate(sales_headers, 1):
        cell = sales_ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
    
    # Sales data
    sales_data = [
        [1001, "C001", 5, 29.99, "2024-01-15"],
        [1002, "C002", 3, 19.99, "2024-01-16"],
        [1003, "C001", 2, 49.99, "2024-01-17"],
        [1004, "C003", 1, 99.99, "2024-01-18"],
        [1005, "C002", 4, 15.99, "2024-01-19"],
        [1006, "C004", 2, 79.99, "2024-01-20"],
        [1007, "C001", 3, 39.99, "2024-01-21"],
        [1008, "C003", 1, 129.99, "2024-01-22"],
    ]
    
    for row, data in enumerate(sales_data, 2):
        for col, value in enumerate(data, 1):
            sales_ws.cell(row=row, column=col, value=value)
    
    # Create Products sheet
    products_ws = wb.create_sheet("Products")
    
    # Products data headers
    products_headers = ["ProductID", "ProductName", "Category", "Supplier", "Cost"]
    for col, header in enumerate(products_headers, 1):
        cell = products_ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
    
    # Products data
    products_data = [
        [1001, "Laptop Pro", "Electronics", "TechCorp", 20.00],
        [1002, "Wireless Mouse", "Electronics", "TechCorp", 12.00],
        [1003, "Gaming Headset", "Electronics", "GameTech", 35.00],
        [1004, "4K Monitor", "Electronics", "DisplayMax", 70.00],
        [1005, "USB Cable", "Accessories", "CableCo", 8.00],
        [1006, "Mechanical Keyboard", "Electronics", "KeyTech", 55.00],
        [1007, "Webcam HD", "Electronics", "VideoPro", 25.00],
        [1008, "Graphics Card", "Electronics", "GPUCorp", 90.00],
    ]
    
    for row, data in enumerate(products_data, 2):
        for col, value in enumerate(data, 1):
            products_ws.cell(row=row, column=col, value=value)
    
    # Create Customers sheet
    customers_ws = wb.create_sheet("Customers")
    
    # Customers data headers
    customers_headers = ["CustomerID", "CustomerName", "Email", "Phone", "City"]
    for col, header in enumerate(customers_headers, 1):
        cell = customers_ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
    
    # Customers data
    customers_data = [
        ["C001", "John Smith", "john.smith@email.com", "555-0101", "New York"],
        ["C002", "Sarah Johnson", "sarah.j@email.com", "555-0102", "Los Angeles"],
        ["C003", "Mike Davis", "mike.davis@email.com", "555-0103", "Chicago"],
        ["C004", "Lisa Wilson", "lisa.w@email.com", "555-0104", "Houston"],
    ]
    
    for row, data in enumerate(customers_data, 2):
        for col, value in enumerate(data, 1):
            customers_ws.cell(row=row, column=col, value=value)
    
    # Save the workbook
    filename = "sample_data.xlsx"
    wb.save(filename)
    print(f"✅ Sample Excel file created: {filename}")
    print("\n📊 File contains 3 sheets:")
    print("  1. Sales - Contains sales transactions")
    print("  2. Products - Contains product information")
    print("  3. Customers - Contains customer information")
    print("\n💡 You can now use this file to test the XLOOKUP tool!")

if __name__ == "__main__":
    create_sample_excel() 